from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func

from app import app
from extensions import db
from models.category import Category
from models.place import Place


# =========================================================
# CONFIGURATION
# =========================================================

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data_import"


IMPORT_JOBS = [
    # =====================================================
    # BEIRUT
    # =====================================================
    {
        "file": "Discover Lebanon Database.xlsx",
        "sheet": "Landmarks & Attractions",
        "header_row": 1,
        "format": "beirut",
        "kind": "activity",
    },
    {
        "file": "Discover Lebanon Database.xlsx",
        "sheet": "Restaurants",
        "header_row": 1,
        "format": "beirut",
        "kind": "restaurant",
    },
    {
        "file": "Discover Lebanon Database.xlsx",
        "sheet": "Coffee Shops",
        "header_row": 1,
        "format": "beirut",
        "kind": "coffee",
    },

    # =====================================================
    # CHOUF - DEEP RESEARCH
    # =====================================================
    {
        "file": "Discover_Lebanon_Chouf_Deep_Research.xlsx",
        "sheet": "Restaurants",
        "header_row": 4,
        "format": "chouf",
        "kind": "restaurant",
    },
    {
        "file": "Discover_Lebanon_Chouf_Deep_Research.xlsx",
        "sheet": "Coffee Shops",
        "header_row": 4,
        "format": "chouf",
        "kind": "coffee",
    },
    {
        "file": "Discover_Lebanon_Chouf_Deep_Research.xlsx",
        "sheet": "Landmarks & Attractions",
        "header_row": 4,
        "format": "chouf",
        "kind": "activity",
    },

    # =====================================================
    # SIDON - DEEP RESEARCH
    # =====================================================
    {
        "file": "Discover_Lebanon_Sidon_Deep_Research.xlsx",
        "sheet": "Restaurants",
        "header_row": 4,
        "format": "sidon",
        "kind": "restaurant",
    },
    {
        "file": "Discover_Lebanon_Sidon_Deep_Research.xlsx",
        "sheet": "Coffee Shops",
        "header_row": 4,
        "format": "sidon",
        "kind": "coffee",
    },
    {
        "file": "Discover_Lebanon_Sidon_Deep_Research.xlsx",
        "sheet": "Landmarks & Attractions",
        "header_row": 4,
        "format": "sidon",
        "kind": "activity",
    },

    # =====================================================
    # TYRE - DEEP RESEARCH
    # =====================================================
    {
        "file": "Discover_Lebanon_Tyre_Deep_Research.xlsx",
        "sheet": "Restaurants",
        "header_row": 4,
        "format": "tyre",
        "kind": "restaurant",
    },
    {
        "file": "Discover_Lebanon_Tyre_Deep_Research.xlsx",
        "sheet": "Coffee Shops",
        "header_row": 4,
        "format": "tyre",
        "kind": "coffee",
    },
    {
        "file": "Discover_Lebanon_Tyre_Deep_Research.xlsx",
        "sheet": "Landmarks & Attractions",
        "header_row": 4,
        "format": "tyre",
        "kind": "activity",
    },

    # =====================================================
    # KESERWAN / JBEIL
    # =====================================================
    {
        "file":
            "Discover_Lebanon_Zone2_Keserwan_Jbeil_Deep_Research.xlsx",
        "sheet": "Food & Drink",
        "header_row": 1,
        "format": "zone2_food",
        "kind": "food",
    },
    {
        "file":
            "Discover_Lebanon_Zone2_Keserwan_Jbeil_Deep_Research.xlsx",
        "sheet": "Activities",
        "header_row": 1,
        "format": "zone2_activity",
        "kind": "activity",
    },

    # =====================================================
    # NABATIEH
    # =====================================================
    {
        "file":
            "Nabatieh_42_Villages_Restaurants_Deep_Research(1).xlsx",
        "sheet": "Restaurants",
        "header_row": 1,
        "format": "nabatieh_restaurant",
        "kind": "restaurant",
    },
    {
        "file":
            "Nabatieh_42_Villages_Coffee_Shops_Deep_Research(1).xlsx",
        "sheet": "Coffee Shops",
        "header_row": 1,
        "format": "nabatieh_coffee",
        "kind": "coffee",
    },
    {
        "file":
            "Nabatieh_42_Villages_Activities_Deep_Research(1).xlsx",
        "sheet": "Activities",
        "header_row": 1,
        "format": "nabatieh_activity",
        "kind": "activity",
    },

    # =====================================================
    # NORTH LEBANON / AKKAR
    # =====================================================
    {
        "file":
            "North_Lebanon_Akkar_Restaurants_Deep_Research(1).xlsx",
        "sheet": "Restaurants",
        "header_row": 1,
        "format": "north_restaurant",
        "kind": "restaurant",
    },
    {
        "file":
            "North_Lebanon_Akkar_Activities_Deep_Research (1).xlsx",
        "sheet": "Activities",
        "header_row": 1,
        "format": "north_activity",
        "kind": "activity",
    },

    # =====================================================
    # ZAHLE / WEST BEKAA / BAALBEK / HERMEL
    # =====================================================
    {
        "file":
            "Zahle_WestBekaa_Baalbek_Hermel_Restaurants_Deep_Research.xlsx",
        "sheet": "Restaurants",
        "header_row": 1,
        "format": "bekaa_restaurant",
        "kind": "restaurant",
    },
    {
        "file":
            "Zahle_WestBekaa_Baalbek_Hermel_CoffeeShops_Deep_Research.xlsx",
        "sheet": "Coffee Shops",
        "header_row": 1,
        "format": "bekaa_coffee",
        "kind": "coffee",
    },
    {
        "file":
            "Zahle_WestBekaa_Baalbek_Hermel_Activities_Deep_Research.xlsx",
        "sheet": "Activities",
        "header_row": 1,
        "format": "bekaa_activity",
        "kind": "activity",
    },
]


# =========================================================
# HELPERS
# =========================================================

def clean(value):
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()

        if not value:
            return None

    return value


def text(value, max_length=None):
    value = clean(value)

    if value is None:
        return None

    value = str(value).strip()

    if max_length:
        return value[:max_length]

    return value


def to_float(value):
    value = clean(value)

    if value is None:
        return None

    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def to_int(value):
    value = clean(value)

    if value is None:
        return None

    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def to_date(value):
    value = clean(value)

    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    value = str(value).strip()

    for fmt in (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(
                value,
                fmt
            ).date()
        except ValueError:
            pass

    return None


def combine_notes(*parts):
    result = []

    for part in parts:
        part = text(part)

        if part:
            result.append(part)

    if not result:
        return None

    return " | ".join(result)


def short_description(description):
    description = text(description)

    if not description:
        return None

    if len(description) <= 300:
        return description

    return description[:297] + "..."


def strip_district(value):
    value = text(value)

    if not value:
        return None

    if value.lower().endswith(" district"):
        return value[:-9].strip()

    return value


# =========================================================
# CATEGORIES
# =========================================================

category_cache = {}


def get_category(name):
    if not name:
        return None

    key = name.strip().lower()

    if key in category_cache:
        return category_cache[key]

    category = Category.query.filter(
        func.lower(Category.name) == key
    ).first()

    if category is None:
        category = Category(
            name=name.strip()
        )

        db.session.add(category)
        db.session.flush()

    category_cache[key] = category

    return category


def activity_frontend_category(
    venue_type,
    category_group=None
):
    combined = (
        f"{venue_type or ''} "
        f"{category_group or ''}"
    ).lower()

    historical_words = [
        "heritage",
        "historical",
        "history",
        "museum",
        "archaeology",
        "archaeological",
        "religious",
        "castle",
        "temple",
        "souk",
        "pilgrimage",
        "cathedral",
        "church",
        "mosque",
        "culture",
    ]

    nature_words = [
        "nature",
        "hiking",
        "waterfall",
        "spring",
        "lake",
        "water",
        "beach",
        "park",
        "geology",
        "outdoor",
        "valley",
        "forest",
        "cedar",
    ]

    if any(
        word in combined
        for word in historical_words
    ):
        return "Historical Sites"

    if any(
        word in combined
        for word in nature_words
    ):
        return "Nature & Outdoors"

    return "Activities & Experiences"


def food_kind(venue_type):
    value = (
        venue_type or ""
    ).lower()

    coffee_words = [
        "cafe",
        "café",
        "coffee",
        "patisserie",
        "dessert",
        "bakery",
    ]

    if any(
        word in value
        for word in coffee_words
    ):
        return "coffee"

    return "restaurant"


def categories_for(
    kind,
    venue_type=None,
    category_group=None
):
    names = []

    if kind == "restaurant":
        names.extend([
            "Restaurant",
            "Food & Gastronomy",
        ])

    elif kind == "coffee":
        names.extend([
            "Coffee Shop",
            "Food & Gastronomy",
        ])

    elif kind == "activity":
        names.append("Activity")

        names.append(
            activity_frontend_category(
                venue_type,
                category_group
            )
        )

    # Add specific research category too.
    if category_group:
        names.append(category_group)

    result = []

    for name in names:
        if (
            name
            and name not in result
        ):
            result.append(name)

    return result


# =========================================================
# REGION NORMALIZATION
# =========================================================

def south_region(locality):
    locality = (
        text(locality) or ""
    ).lower()

    chouf_names = [
        "deir el qamar",
        "beiteddine",
        "baaqline",
        "baakleen",
        "dmit",
        "chouf",
    ]

    if any(
        item in locality
        for item in chouf_names
    ):
        return "Chouf"

    if "sidon" in locality:
        return "Sidon"

    if "tyre" in locality:
        return "Tyre"

    if "naqoura" in locality:
        return "Naqoura"

    return text(locality)


# =========================================================
# WORKBOOK READING
# =========================================================

def read_rows(
    filepath,
    sheet_name,
    header_row
):
    workbook = load_workbook(
        filepath,
        read_only=True,
        data_only=True
    )

    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError(
            f"Sheet '{sheet_name}' not found "
            f"in {filepath.name}"
        )

    sheet = workbook[sheet_name]

    headers = []

    for cell in sheet[header_row]:
        headers.append(
            text(cell.value)
        )

    records = []

    for values in sheet.iter_rows(
        min_row=header_row + 1,
        values_only=True
    ):
        record = {}

        has_data = False

        for header, value in zip(
            headers,
            values
        ):
            if not header:
                continue

            value = clean(value)

            record[header] = value

            if value is not None:
                has_data = True

        if has_data:
            records.append(record)

    workbook.close()

    return records


# =========================================================
# RECORD NORMALIZATION
# =========================================================

def normalize_beirut(row, kind):
    name = row.get("Name")

    if not name:
        return None

    venue_type = row.get("Type")
    region = row.get("Region")
    locality = row.get("District/Area")
    description = row.get("Description")

    return {
        "name": text(name, 150),
        "short_description":
            short_description(description),
        "description": text(description),

        "venue_type":
            text(venue_type, 150),

        "category_group": None,

        "city":
            text(locality, 100),

        "region":
            text(region, 100),

        "address":
            text(locality, 255),

        "opening_hours": None,
        "phone": None,

        "status":
            "Research dataset",

        "rating": None,
        "review_count": None,

        "public_access": None,
        "confidence": None,
        "locality_match": None,
        "evidence_source": None,
        "source_url": None,
        "map_url": None,

        "notes": combine_notes(
            (
                f"Price: "
                f"{row.get('Price Range')}"
                if row.get("Price Range")
                else None
            ),
            row.get("Tips / Notes"),
        ),

        "last_checked": None,

        "_categories":
            categories_for(
                kind,
                venue_type
            ),
    }


def normalize_zone4(row, kind):
    name = row.get("Name")

    if not name:
        return None

    locality = row.get("Town / Area")
    description = row.get(
        "Notes / Highlights"
    )

    venue_type = (
        row.get("Category")
        or (
            "Restaurant"
            if kind == "restaurant"
            else "Coffee Shop"
            if kind == "coffee"
            else "Attraction"
        )
    )

    return {
        "name": text(name, 150),

        "short_description":
            short_description(description),

        "description":
            text(description),

        "venue_type":
            text(venue_type, 150),

        "category_group": None,

        "city":
            text(locality, 100),

        "region":
            text(
                south_region(locality),
                100
            ),

        "address":
            text(
                row.get("Address"),
                255
            ),

        "opening_hours":
            text(
                row.get(
                    "Opening Hours (Sample)"
                ),
                200
            ),

        "phone":
            text(
                row.get("Phone Number"),
                50
            ),

        "status":
            "Research dataset",

        "rating":
            to_float(
                row.get("Rating (Google)")
            ),

        "review_count":
            to_int(
                row.get("Reviews Count")
            ),

        "public_access": None,
        "confidence": None,
        "locality_match": None,
        "evidence_source": None,
        "source_url": None,
        "map_url": None,

        "notes": combine_notes(
            (
                f"Price: "
                f"{row.get('Price Level')}"
                if row.get("Price Level")
                else None
            ),
            description,
        ),

        "last_checked": None,

        "_categories":
            categories_for(
                kind,
                venue_type
            ),
    }


def normalize_chouf(row, kind):
    name = row.get("Name")

    if not name:
        return None

    locality = row.get("Town / Area")
    description = row.get(
        "Notes / Highlights"
    )

    venue_type = (
        row.get("Category")
        or (
            "Restaurant"
            if kind == "restaurant"
            else "Coffee Shop"
            if kind == "coffee"
            else "Attraction"
        )
    )

    return {
        "name": text(name, 150),

        "short_description":
            short_description(description),

        "description":
            text(description),

        "venue_type":
            text(venue_type, 150),

        "category_group": None,

        # Village / locality inside Chouf
        "city":
            text(locality, 100),

        # Every record in this workbook belongs to Chouf
        "region": "Chouf",

        "address":
            text(
                row.get("Address"),
                255
            ),

        "opening_hours":
            text(
                row.get(
                    "Opening Hours (Sample)"
                ),
                200
            ),

        "phone":
            text(
                row.get("Phone Number"),
                50
            ),

        "status":
            text(
                row.get("Status"),
                100
            )
            or "Research dataset",

        "rating":
            to_float(
                row.get("Rating (Google)")
            ),

        "review_count":
            to_int(
                row.get("Reviews Count")
            ),

        "public_access": None,
        "confidence": None,
        "locality_match": None,

        "evidence_source":
            "Chouf deep research",

        "source_url":
            text(
                row.get("Source URL")
            ),

        "map_url":
            text(
                row.get(
                    "Map / Search Link"
                )
            ),

        "notes":
            combine_notes(
                (
                    f"Price: "
                    f"{row.get('Price Level')}"
                    if row.get("Price Level")
                    else None
                ),
                description,
            ),

        "last_checked":
            to_date(
                row.get("Last Checked")
            ),

        "_categories":
            categories_for(
                kind,
                venue_type
            ),
    }



def normalize_sidon(row, kind):
    name = row.get("Name")

    if not name:
        return None

    locality = row.get("Town / Area")
    description = row.get(
        "Notes / Highlights"
    )

    venue_type = (
        row.get("Category")
        or (
            "Restaurant"
            if kind == "restaurant"
            else "Coffee Shop"
            if kind == "coffee"
            else "Attraction"
        )
    )

    return {
        "name": text(name, 150),

        "short_description":
            short_description(description),

        "description":
            text(description),

        "venue_type":
            text(venue_type, 150),

        "category_group": None,

        "city":
            text(locality, 100),

        "region": "Sidon",

        "address":
            text(
                row.get("Address"),
                255
            ),

        "opening_hours":
            text(
                row.get(
                    "Opening Hours (Sample)"
                ),
                200
            ),

        "phone":
            text(
                row.get("Phone Number"),
                50
            ),

        "status":
            text(
                row.get("Status"),
                100
            )
            or "Research dataset",

        "rating":
            to_float(
                row.get("Rating (Google)")
            ),

        "review_count":
            to_int(
                row.get("Reviews Count")
            ),

        "public_access": None,
        "confidence": None,
        "locality_match": None,

        "evidence_source":
            "Sidon deep research",

        "source_url":
            text(
                row.get("Source URL")
            ),

        "map_url":
            text(
                row.get(
                    "Map / Search Link"
                )
            ),

        "notes":
            combine_notes(
                (
                    f"Price: "
                    f"{row.get('Price Level')}"
                    if row.get("Price Level")
                    else None
                ),
                description,
            ),

        "last_checked":
            to_date(
                row.get("Last Checked")
            ),

        "_categories":
            categories_for(
                kind,
                venue_type
            ),
    }


def normalize_tyre(row, kind):
    name = row.get("Name")

    if not name:
        return None

    locality = row.get("Town / Area")
    description = row.get(
        "Notes / Highlights"
    )

    venue_type = (
        row.get("Category")
        or (
            "Restaurant"
            if kind == "restaurant"
            else "Coffee Shop"
            if kind == "coffee"
            else "Attraction"
        )
    )

    return {
        "name": text(name, 150),

        "short_description":
            short_description(description),

        "description":
            text(description),

        "venue_type":
            text(venue_type, 150),

        "category_group": None,

        "city":
            text(locality, 100),

        "region": "Tyre",

        "address":
            text(
                row.get("Address"),
                255
            ),

        "opening_hours":
            text(
                row.get(
                    "Opening Hours (Sample)"
                ),
                200
            ),

        "phone":
            text(
                row.get("Phone Number"),
                50
            ),

        "status":
            text(
                row.get("Status"),
                100
            )
            or "Research dataset",

        "rating":
            to_float(
                row.get("Rating (Google)")
            ),

        "review_count":
            to_int(
                row.get("Reviews Count")
            ),

        "public_access": None,
        "confidence": None,
        "locality_match": None,

        "evidence_source":
            "Tyre deep research",

        "source_url":
            text(
                row.get("Source URL")
            ),

        "map_url":
            text(
                row.get(
                    "Map / Search Link"
                )
            ),

        "notes":
            combine_notes(
                (
                    f"Price: "
                    f"{row.get('Price Level')}"
                    if row.get("Price Level")
                    else None
                ),
                description,
            ),

        "last_checked":
            to_date(
                row.get("Last Checked")
            ),

        "_categories":
            categories_for(
                kind,
                venue_type
            ),
    }

def normalize_zone2_food(row):
    name = row.get("Venue Name")

    if not name:
        return None

    venue_type = row.get(
        "Type / Cuisine"
    )

    kind = food_kind(venue_type)

    return {
        "name":
            text(name, 150),

        "short_description":
            short_description(
                row.get("Notes")
            ),

        "description":
            text(row.get("Notes")),

        "venue_type":
            text(venue_type, 150),

        "category_group": None,

        "city":
            text(
                row.get("Exact Locality"),
                100
            ),

        "region":
            text(
                row.get("Area"),
                100
            ),

        "address":
            text(
                row.get("Exact Locality"),
                255
            ),

        "opening_hours": None,
        "phone": None,

        "status":
            text(
                row.get("Status"),
                100
            ),

        "rating":
            to_float(
                row.get("Rating")
            ),

        "review_count":
            to_int(
                row.get("Reviews")
            ),

        "public_access": None,

        "confidence":
            text(
                row.get("Confidence"),
                50
            ),

        "locality_match": None,

        "evidence_source":
            text(
                row.get("Evidence"),
                255
            ),

        "source_url":
            text(
                row.get("Source URL")
            ),

        "map_url":
            text(
                row.get(
                    "Map / Search Link"
                )
            ),

        "notes":
            text(row.get("Notes")),

        "last_checked":
            to_date(
                row.get("Last Checked")
            ),

        "_categories":
            categories_for(
                kind,
                venue_type
            ),
    }


def normalize_zone2_activity(row):
    name = row.get(
        "Activity / Place"
    )

    if not name:
        return None

    venue_type = row.get("Category")

    return {
        "name":
            text(name, 150),

        "short_description":
            short_description(
                row.get(
                    "What You Can Do"
                )
            ),

        "description":
            text(
                row.get(
                    "What You Can Do"
                )
            ),

        "venue_type":
            text(venue_type, 150),

        "category_group": None,

        "city":
            text(
                row.get("Exact Locality"),
                100
            ),

        "region":
            text(
                row.get("Area"),
                100
            ),

        "address":
            text(
                row.get("Exact Locality"),
                255
            ),

        "opening_hours":
            text(
                row.get("Season"),
                200
            ),

        "phone": None,
        "status": "Research dataset",

        "rating":
            to_float(
                row.get("Rating")
            ),

        "review_count":
            to_int(
                row.get("Reviews")
            ),

        "public_access": None,

        "confidence":
            text(
                row.get("Confidence"),
                50
            ),

        "locality_match": None,
        "evidence_source": None,

        "source_url":
            text(
                row.get("Source URL")
            ),

        "map_url":
            text(
                row.get(
                    "Map / Search Link"
                )
            ),

        "notes":
            text(row.get("Notes")),

        "last_checked":
            to_date(
                row.get("Last Checked")
            ),

        "_categories":
            categories_for(
                "activity",
                venue_type
            ),
    }


def normalize_nabatieh_restaurant(row):
    name = row.get("Business Name")

    if not name:
        return None

    venue_type = row.get(
        "Venue Type"
    )

    return {
        "name":
            text(name, 150),

        "short_description":
            short_description(
                row.get("Notes")
            ),

        "description":
            text(row.get("Notes")),

        "venue_type":
            text(venue_type, 150),

        "category_group": None,

        "city":
            text(
                row.get(
                    "Canonical Locality"
                ),
                100
            ),

        "region": "Nabatieh",

        "address":
            text(
                row.get(
                    "Address / Area"
                ),
                255
            ),

        "opening_hours":
            text(
                row.get("Hours"),
                200
            ),

        "phone":
            text(
                row.get("Phone"),
                50
            ),

        "status":
            text(
                row.get("Status"),
                100
            ),

        "rating":
            to_float(
                row.get("Rating")
            ),

        "review_count":
            to_int(
                row.get("Reviews")
            ),

        "public_access": None,

        "confidence":
            text(
                row.get("Confidence"),
                50
            ),

        "locality_match":
            text(
                row.get(
                    "Locality Match"
                ),
                150
            ),

        "evidence_source":
            text(
                row.get(
                    "Evidence Source"
                ),
                255
            ),

        "source_url":
            text(
                row.get("Source URL")
            ),

        "map_url":
            text(
                row.get(
                    "Map / Search Link"
                )
            ),

        "notes":
            text(row.get("Notes")),

        "last_checked":
            to_date(
                row.get("Last Checked")
            ),

        "_categories":
            categories_for(
                "restaurant",
                venue_type
            ),
    }


def normalize_nabatieh_coffee(row):
    name = row.get("Business Name")

    if not name:
        return None

    venue_type = row.get("Type")

    return {
        "name":
            text(name, 150),

        "short_description":
            short_description(
                row.get("Notes")
            ),

        "description":
            text(row.get("Notes")),

        "venue_type":
            text(venue_type, 150),

        "category_group": None,

        "city":
            text(
                row.get(
                    "Canonical Locality"
                ),
                100
            ),

        "region": "Nabatieh",

        "address":
            text(
                row.get(
                    "Address / Area"
                ),
                255
            ),

        "opening_hours":
            text(
                row.get("Hours"),
                200
            ),

        "phone":
            text(
                row.get("Phone"),
                50
            ),

        "status":
            text(
                row.get("Status"),
                100
            ),

        "rating":
            to_float(
                row.get("Rating")
            ),

        "review_count":
            to_int(
                row.get("Reviews")
            ),

        "public_access": None,

        "confidence":
            text(
                row.get("Confidence"),
                50
            ),

        "locality_match": None,

        "evidence_source":
            text(
                row.get(
                    "Evidence Source"
                ),
                255
            ),

        "source_url":
            text(
                row.get("Source URL")
            ),

        "map_url":
            text(
                row.get(
                    "Map / Search Link"
                )
            ),

        "notes":
            text(row.get("Notes")),

        "last_checked":
            to_date(
                row.get("Last Checked")
            ),

        "_categories":
            categories_for(
                "coffee",
                venue_type
            ),
    }


def normalize_nabatieh_activity(row):
    name = row.get(
        "Activity / Place"
    )

    if not name:
        return None

    venue_type = row.get("Category")
    group = row.get(
        "Category Group"
    )

    return {
        "name":
            text(name, 150),

        "short_description":
            short_description(
                row.get(
                    "What You Can Do"
                )
            ),

        "description":
            text(
                row.get(
                    "What You Can Do"
                )
            ),

        "venue_type":
            text(venue_type, 150),

        "category_group":
            text(group, 100),

        "city":
            text(
                row.get(
                    "Canonical Locality"
                ),
                100
            ),

        "region": "Nabatieh",

        "address":
            text(
                row.get(
                    "Address / Area"
                ),
                255
            ),

        "opening_hours":
            text(
                row.get("Hours"),
                200
            ),

        "phone":
            text(
                row.get("Phone"),
                50
            ),

        "status":
            text(
                row.get("Status"),
                100
            ),

        "rating":
            to_float(
                row.get("Rating")
            ),

        "review_count":
            to_int(
                row.get("Reviews")
            ),

        "public_access":
            text(
                row.get(
                    "Public Access"
                ),
                100
            ),

        "confidence":
            text(
                row.get("Confidence"),
                50
            ),

        "locality_match":
            text(
                row.get(
                    "Locality Match"
                ),
                150
            ),

        "evidence_source":
            text(
                row.get(
                    "Evidence Source"
                ),
                255
            ),

        "source_url":
            text(
                row.get("Source URL")
            ),

        "map_url":
            text(
                row.get(
                    "Map / Search Link"
                )
            ),

        "notes":
            text(row.get("Notes")),

        "last_checked":
            to_date(
                row.get("Last Checked")
            ),

        "_categories":
            categories_for(
                "activity",
                venue_type,
                group
            ),
    }


def normalize_north_restaurant(row):
    name = row.get(
        "Restaurant Name"
    )

    if not name:
        return None

    venue_type = row.get(
        "Type / Cuisine"
    )

    return {
        "name":
            text(name, 150),

        "short_description":
            short_description(
                row.get("Notes")
            ),

        "description":
            text(row.get("Notes")),

        "venue_type":
            text(venue_type, 150),

        "category_group": None,

        "city":
            text(
                row.get("Exact Locality"),
                100
            ),

        "region":
            text(
                row.get("Area"),
                100
            ),

        "address":
            text(
                row.get("Exact Locality"),
                255
            ),

        "opening_hours":
            text(
                row.get("Hours"),
                200
            ),

        "phone":
            text(
                row.get("Phone"),
                50
            ),

        "status":
            text(
                row.get("Status"),
                100
            ),

        "rating":
            to_float(
                row.get("Rating")
            ),

        "review_count":
            to_int(
                row.get("Reviews")
            ),

        "public_access": None,

        "confidence":
            text(
                row.get("Confidence"),
                50
            ),

        "locality_match": None,

        "evidence_source":
            text(
                row.get("Evidence"),
                255
            ),

        "source_url":
            text(
                row.get("Source URL")
            ),

        "map_url":
            text(
                row.get(
                    "Map / Search Link"
                )
            ),

        "notes":
            text(row.get("Notes")),

        "last_checked":
            to_date(
                row.get("Last Checked")
            ),

        "_categories":
            categories_for(
                "restaurant",
                venue_type
            ),
    }


def normalize_north_activity(row):
    name = row.get(
        "Activity / Place"
    )

    if not name:
        return None

    venue_type = row.get("Category")

    description = row.get(
        "What You Can Do"
    )

    return {
        "name":
            text(name, 150),

        "short_description":
            short_description(
                description
            ),

        "description":
            text(description),

        "venue_type":
            text(venue_type, 150),

        "category_group": None,

        "city":
            text(
                row.get("Exact Locality"),
                100
            ),

        "region":
            text(
                row.get("Area"),
                100
            ),

        "address":
            text(
                row.get(
                    "Address / Area"
                ),
                255
            ),

        "opening_hours":
            text(
                row.get(
                    "Hours / Seasonality"
                ),
                200
            ),

        "phone":
            text(
                row.get("Phone"),
                50
            ),

        "status":
            text(
                row.get("Status"),
                100
            ),

        "rating":
            to_float(
                row.get("Rating")
            ),

        "review_count":
            to_int(
                row.get("Reviews")
            ),

        "public_access": None,

        "confidence":
            text(
                row.get("Confidence"),
                50
            ),

        "locality_match": None,

        "evidence_source":
            text(
                row.get(
                    "Evidence Source"
                ),
                255
            ),

        "source_url":
            text(
                row.get("Source URL")
            ),

        "map_url":
            text(
                row.get(
                    "Map / Search Link"
                )
            ),

        "notes":
            text(row.get("Notes")),

        "last_checked":
            to_date(
                row.get("Last Checked")
            ),

        "_categories":
            categories_for(
                "activity",
                venue_type
            ),
    }


def normalize_bekaa_restaurant(row):
    name = row.get(
        "Restaurant Name"
    )

    if not name:
        return None

    venue_type = row.get(
        "Cuisine / Type"
    )

    description = row.get(
        "Description"
    )

    return {
        "name":
            text(name, 150),

        "short_description":
            short_description(
                description
            ),

        "description":
            text(description),

        "venue_type":
            text(venue_type, 150),

        "category_group": None,

        "city":
            text(
                row.get("Locality"),
                100
            ),

        "region":
            text(
                strip_district(
                    row.get("District")
                ),
                100
            ),

        "address":
            text(
                row.get(
                    "Address / Area"
                ),
                255
            ),

        "opening_hours":
            text(
                row.get("Hours"),
                200
            ),

        "phone":
            text(
                row.get("Phone"),
                50
            ),

        "status":
            text(
                row.get("Status"),
                100
            ),

        "rating":
            to_float(
                row.get("Rating")
            ),

        "review_count":
            to_int(
                row.get("Reviews")
            ),

        "public_access": None,
        "confidence": None,
        "locality_match": None,

        "evidence_source":
            text(
                row.get(
                    "Evidence Source"
                ),
                255
            ),

        "source_url":
            text(
                row.get("Source URL")
            ),

        "map_url": None,

        "notes": None,

        "last_checked":
            to_date(
                row.get("Last Checked")
            ),

        "_categories":
            categories_for(
                "restaurant",
                venue_type
            ),
    }


def normalize_bekaa_coffee(row):
    name = row.get(
        "Coffee Shop Name"
    )

    if not name:
        return None

    description = row.get(
        "Description"
    )

    return {
        "name":
            text(name, 150),

        "short_description":
            short_description(
                description
            ),

        "description":
            text(description),

        "venue_type": "Coffee Shop",
        "category_group": None,

        "city":
            text(
                row.get("Locality"),
                100
            ),

        "region":
            text(
                strip_district(
                    row.get("District")
                ),
                100
            ),

        "address":
            text(
                row.get(
                    "Address / Area"
                ),
                255
            ),

        "opening_hours":
            text(
                row.get("Hours"),
                200
            ),

        "phone":
            text(
                row.get("Phone"),
                50
            ),

        "status":
            text(
                row.get("Status"),
                100
            ),

        "rating":
            to_float(
                row.get("Rating")
            ),

        "review_count":
            to_int(
                row.get("Reviews")
            ),

        "public_access": None,
        "confidence": None,
        "locality_match": None,

        "evidence_source":
            text(
                row.get(
                    "Evidence Source"
                ),
                255
            ),

        "source_url":
            text(
                row.get("Source URL")
            ),

        "map_url": None,

        "notes": None,

        "last_checked":
            to_date(
                row.get("Last Checked")
            ),

        "_categories":
            categories_for(
                "coffee",
                "Coffee Shop"
            ),
    }


def normalize_bekaa_activity(row):
    name = row.get(
        "Activity / Place"
    )

    if not name:
        return None

    venue_type = row.get(
        "Category"
    )

    group = row.get(
        "Category Group"
    )

    description = row.get(
        "What You Can Do"
    )

    return {
        "name":
            text(name, 150),

        "short_description":
            short_description(
                description
            ),

        "description":
            text(description),

        "venue_type":
            text(venue_type, 150),

        "category_group":
            text(group, 100),

        "city":
            text(
                row.get("Locality"),
                100
            ),

        "region":
            text(
                strip_district(
                    row.get("District")
                ),
                100
            ),

        "address":
            text(
                row.get(
                    "Address / Area"
                ),
                255
            ),

        "opening_hours":
            text(
                row.get("Hours"),
                200
            ),

        "phone":
            text(
                row.get("Phone"),
                50
            ),

        "status":
            text(
                row.get("Status"),
                100
            ),

        "rating":
            to_float(
                row.get("Rating")
            ),

        "review_count":
            to_int(
                row.get("Reviews")
            ),

        "public_access": None,
        "confidence": None,
        "locality_match": None,

        "evidence_source":
            text(
                row.get(
                    "Evidence Source"
                ),
                255
            ),

        "source_url":
            text(
                row.get("Source URL")
            ),

        "map_url": None,

        "notes":
            text(row.get("Notes")),

        "last_checked":
            to_date(
                row.get("Last Checked")
            ),

        "_categories":
            categories_for(
                "activity",
                venue_type,
                group
            ),
    }


# =========================================================
# NORMALIZER DISPATCH
# =========================================================

def normalize_record(
    row,
    job
):
    format_name = job["format"]
    kind = job["kind"]

    if format_name == "beirut":
        return normalize_beirut(
            row,
            kind
        )

    if format_name == "zone4":
        return normalize_zone4(
            row,
            kind
        )

    if format_name == "chouf":
        return normalize_chouf(
            row,
            kind
        )

    if format_name == "sidon":
        return normalize_sidon(
            row,
            kind
        )

    if format_name == "tyre":
        return normalize_tyre(
            row,
            kind
        )

    if format_name == "zone2_food":
        return normalize_zone2_food(
            row
        )

    if format_name == "zone2_activity":
        return normalize_zone2_activity(
            row
        )

    if format_name == \
            "nabatieh_restaurant":
        return \
            normalize_nabatieh_restaurant(
                row
            )

    if format_name == \
            "nabatieh_coffee":
        return normalize_nabatieh_coffee(
            row
        )

    if format_name == \
            "nabatieh_activity":
        return normalize_nabatieh_activity(
            row
        )

    if format_name == \
            "north_restaurant":
        return normalize_north_restaurant(
            row
        )

    if format_name == \
            "north_activity":
        return normalize_north_activity(
            row
        )

    if format_name == \
            "bekaa_restaurant":
        return normalize_bekaa_restaurant(
            row
        )

    if format_name == \
            "bekaa_coffee":
        return normalize_bekaa_coffee(
            row
        )

    if format_name == \
            "bekaa_activity":
        return normalize_bekaa_activity(
            row
        )

    raise ValueError(
        f"Unknown format: {format_name}"
    )


# =========================================================
# DATABASE UPSERT
# =========================================================

PLACE_FIELDS = [
    "name",
    "short_description",
    "description",
    "venue_type",
    "category_group",
    "city",
    "region",
    "address",
    "opening_hours",
    "phone",
    "status",
    "rating",
    "review_count",
    "public_access",
    "confidence",
    "locality_match",
    "evidence_source",
    "source_url",
    "map_url",
    "notes",
    "last_checked",
]


def find_existing(
    name,
    city,
    region
):
    query = Place.query.filter(
        func.lower(Place.name)
        == name.lower()
    )

    if city:
        query = query.filter(
            func.lower(
                func.coalesce(
                    Place.city,
                    ""
                )
            )
            == city.lower()
        )

    if region:
        query = query.filter(
            func.lower(
                func.coalesce(
                    Place.region,
                    ""
                )
            )
            == region.lower()
        )

    return query.first()

def upsert_place(data):
    name = data.get("name")

    if not name:
        return "skipped"

    city = data.get("city")
    region = data.get("region")

    # Prevent previously identified Tyre duplicates
    tyre_duplicate_locations = {
        # First four duplicates
        ("al fakhoury restaurant", "old port"): "Tyre - Old Port",
        ("bolbol mina resto", "old port"): "Tyre - Mina / Waterfront",
        ("le phenicien sea food restaurant", "old port"): "Tyre - Old Port",
        ("restaurant tony", "old port"): "Tyre - Old Port",

        # Seven additional duplicates
        ("baroud restaurant", "old souk"): "Tyre - Old Souk",
        ("tyre coast nature reserve", "southern tyre coast"): "Tyre - South Coast",
        ("yam seafood resto", "cheikh abbas nasser"): "Tyre - Western Sea",
        ("kahwet al balad", "al kharab street"): "Tyre - Corniche",
        ("bamboo restaurant", "borj ech chemali"): "Borj Ech Chemali / Tyre Area",
        ("tyre hippodrome", "al-bass"): "Tyre - Al-Bass",
        ("fiora patisserie artisanale", "jal el bahr"): "Jal El-Bahr",
    }

    if region and city and region.strip().lower() == "tyre":
        key = (
            name.strip().lower(),
            city.strip().lower()
        )

        canonical_city = tyre_duplicate_locations.get(key)

        if canonical_city:
            canonical_place = find_existing(
                name,
                canonical_city,
                region
            )

            if canonical_place is None:
                raise RuntimeError(
                    f"Canonical Tyre place missing: {name!r} / "
                    f"{canonical_city!r}. Import paused."
                )

            def digits(value):
                return "".join(
                    char for char in (value or "")
                    if char.isdigit()
                )

            incoming_phone = digits(data.get("phone"))
            canonical_phone = digits(canonical_place.phone)

            same_phone = (
                bool(incoming_phone)
                and incoming_phone == canonical_phone
            )

            no_phone_attraction = (
                key in {
                    (
                        "tyre coast nature reserve",
                        "southern tyre coast"
                    ),
                    (
                        "tyre hippodrome",
                        "al-bass"
                    ),
                }
                and not incoming_phone
            )

            if same_phone or no_phone_attraction:

                # Preserve categories from both records
                added_category = False

                for category_name in data.get("_categories", []):
                    category = get_category(category_name)

                    if (
                        category
                        and category not in canonical_place.categories
                    ):
                        canonical_place.categories.append(category)
                        added_category = True

                # Do not overwrite detailed information
                # with less complete duplicate research rows
                return "updated" if added_category else "skipped"

            raise RuntimeError(
                f"Possible distinct Tyre place / phone mismatch: "
                f"{name!r}, incoming city={city!r}. "
                "Review before importing."
            )

    # Normal import for all other places
    place = find_existing(
        name,
        city,
        region
    )

    created = False

    if place is None:
        place = Place(name=name)
        db.session.add(place)
        created = True

    # Update the place fields
    for field in PLACE_FIELDS:
        value = data.get(field)

        if value is not None:
            setattr(place, field, value)

    # Add categories without duplicates
    for category_name in data.get("_categories", []):
        category = get_category(category_name)

        if (
            category
            and category not in place.categories
        ):
            place.categories.append(category)

    return "added" if created else "updated"

# =========================================================
# IMPORT
# =========================================================

def import_all():
    if not DATA_DIR.exists():
        print(
            f"❌ Folder not found: "
            f"{DATA_DIR}"
        )
        return

    totals = {
        "added": 0,
        "updated": 0,
        "skipped": 0,
        "files_missing": 0,
    }

    print()
    print(
        "======================================"
    )
    print(
        " DISCOVER LEBANON DATA IMPORT"
    )
    print(
        "======================================"
    )
    print()

    for job in IMPORT_JOBS:
        filepath = (
            DATA_DIR / job["file"]
        )

        if not filepath.exists():
            print(
                f"⚠ Missing: "
                f"{job['file']}"
            )

            totals[
                "files_missing"
            ] += 1

            continue

        print(
            f"📄 {job['file']}"
        )

        print(
            f"   Sheet: "
            f"{job['sheet']}"
        )

        try:
            rows = read_rows(
                filepath,
                job["sheet"],
                job["header_row"]
            )

            job_added = 0
            job_updated = 0
            job_skipped = 0

            for row in rows:
                data = normalize_record(
                    row,
                    job
                )

                if not data:
                    job_skipped += 1
                    continue

                result = upsert_place(
                    data
                )

                if result == "added":
                    job_added += 1

                elif result == "updated":
                    job_updated += 1

                else:
                    job_skipped += 1

            db.session.commit()

            totals["added"] += (
                job_added
            )

            totals["updated"] += (
                job_updated
            )

            totals["skipped"] += (
                job_skipped
            )

            print(
                f"   ✅ Added: "
                f"{job_added}"
            )

            print(
                f"   🔄 Updated: "
                f"{job_updated}"
            )

            print(
                f"   ⏭ Skipped: "
                f"{job_skipped}"
            )

        except Exception as exc:
            db.session.rollback()

            print(
                f"   ❌ ERROR: {exc}"
            )

        print()

    print(
        "======================================"
    )

    print(
        " IMPORT COMPLETE"
    )

    print(
        "======================================"
    )

    print(
        f"Added: "
        f"{totals['added']}"
    )

    print(
        f"Updated: "
        f"{totals['updated']}"
    )

    print(
        f"Skipped: "
        f"{totals['skipped']}"
    )

    print(
        f"Missing files: "
        f"{totals['files_missing']}"
    )

    print(
        f"Total places in database: "
        f"{Place.query.count()}"
    )

    print(
        f"Total categories: "
        f"{Category.query.count()}"
    )

    print()


if __name__ == "__main__":
    with app.app_context():
        import_all()